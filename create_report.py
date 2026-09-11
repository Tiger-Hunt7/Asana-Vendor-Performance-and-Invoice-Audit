from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Asana_Vendor_Performance_and_Invoice_Audit_Project_Report.docx"

NAVY = RGBColor(23, 54, 93)
BLUE = "2F75B5"
PALE = "EAF3F8"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(25)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = NAVY
for name in ("Heading 1", "Heading 2"):
    styles[name].font.name = "Aptos Display"
    styles[name].font.color.rgb = NAVY

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Asana Vendor Performance\nand Invoice Audit")
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Management-requested payment-dispute analysis")
r.bold = True
r.font.size = Pt(13)
r.font.color.rgb = RGBColor(47, 117, 181)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("PUBLIC PORTFOLIO VERSION")
r.bold = True
r.font.size = Pt(9)
r.font.color.rgb = RGBColor(127, 96, 0)

doc.add_heading("Executive finding", level=1)
p = doc.add_paragraph()
p.add_run("The vendor claimed that more than $40,000 remained unpaid. ").bold = True
p.add_run(
    "The evidence review found $0 payable for completed work remained unpaid. "
    "The $23,685 classified as denied related to work that was not performed or did not satisfy documented requirements; it was therefore not an unpaid obligation."
)

table = doc.add_table(rows=1, cols=2)
table.style = "Table Grid"
for i, text in enumerate(("Key measure", "Finding")):
    cell = table.rows[0].cells[i]
    cell.text = text
    shade(cell, BLUE)
    for run in cell.paragraphs[0].runs:
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
rows = [
    ("Vendor assertion", "More than $40,000 unpaid"),
    ("Payable unpaid amount found", "$0"),
    ("Itemized reported paid", "$139,948"),
    ("Denied — work not performed / requirements unmet", "$23,685"),
    ("Itemized total reviewed", "$163,633"),
]
for label, value in rows:
    cells = table.add_row().cells
    cells[0].text, cells[1].text = label, value
    for cell in cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    shade(cells[0], PALE)

doc.add_heading("Project purpose", level=1)
doc.add_paragraph(
    "Management requested a defensible review of a vendor payment dispute. The project organized exported Asana task history, inspection outcomes, completion records, service periods, and itemized financial classifications into a traceable analytical workbook. The goal was to test the vendor's assertion, document the organization's position, and identify records requiring reconciliation."
)

doc.add_heading("Approach", level=1)
for text in [
    "Standardized the exported Asana task data and separated dated from undated activity.",
    "Measured completion, open work, failed inspections, and status gaps across 147 tasks.",
    "Built an itemized 91-row financial ledger distinguishing reported paid from denied work.",
    "Compared itemized amounts with previously published summaries and documented discrepancies.",
    "Preserved an audit trail while identifying which claims still required confirmation in the authoritative financial system.",
]:
    doc.add_paragraph(text, style="List Bullet")

doc.add_heading("Operational and data-quality findings", level=1)
table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
for i, text in enumerate(("Measure", "Result", "Meaning")):
    table.rows[0].cells[i].text = text
    shade(table.rows[0].cells[i], BLUE)
    for run in table.rows[0].cells[i].paragraphs[0].runs:
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
findings = [
    ("Asana tasks", "147", "Portfolio snapshot reviewed"),
    ("Completed / open", "124 / 23", "Operational workload status"),
    ("Failed inspection", "45", "Documented performance concern"),
    ("Undated tasks", "21", "Cannot be assigned reliably to a service month"),
    ("Tasks without status", "57", "Outcome cannot be inferred"),
    ("Paid-summary variance", "$12,208", "Older summary exceeded itemized total"),
    ("Denied-summary variance", "$1,847", "Older summary exceeded itemized total"),
]
for row in findings:
    cells = table.add_row().cells
    for i, value in enumerate(row):
        cells[i].text = value
    shade(cells[0], PALE)

doc.add_heading("Conclusion and control point", level=1)
doc.add_paragraph(
    "The analysis supported the organization's position with detailed task and financial evidence: no payable balance was identified for completed work. The review also demonstrated why project-management data should not be treated as a payment ledger. Asana was reliable for operational dates, completion, sections, and inspection history, while final payment certification required reconciliation to the authoritative financial system and underlying invoice documentation."
)

doc.add_heading("Public portfolio data notice", level=1)
doc.add_paragraph(
    "This repository is a fictionalized portfolio version. Vendor, organization, employee, location, task, invoice, link, and identifying narrative details were replaced. The numerical findings and analytical structure were retained at the user's direction to demonstrate the work accurately without identifying the original parties."
)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("Asana Vendor Performance and Invoice Audit | Fictionalized portfolio version")
run.font.size = Pt(8)
run.font.color.rgb = RGBColor(100, 100, 100)

doc.save(OUT)
print(OUT)
